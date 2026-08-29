from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from .base import ImportBookError, ImportResult, ParagraphData, ProgressCallback, SectionData, fallback_title


class DOCXImporter:
    def import_book(self, path: Path, progress: ProgressCallback) -> ImportResult:
        try:
            document = Document(path)
        except PackageNotFoundError as exc:
            raise ImportBookError("DOCX 文件损坏或受保护，无法读取。") from exc
        title = (document.core_properties.title or fallback_title(path)).strip()
        author = (document.core_properties.author or "").strip() or None
        sections: list[SectionData] = []
        current_title = "正文"
        current: list[ParagraphData] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                if current:
                    sections.append(SectionData(current_title, current))
                    current = []
                current_title = value
            else:
                current.append(ParagraphData(value))
        if current:
            sections.append(SectionData(current_title, current))
        if not sections:
            sections = [SectionData("正文", [ParagraphData("[文档中没有可阅读文字]")])]
        progress(len(sections), len(sections), "DOCX 整理完成")
        return ImportResult(
            title=title,
            author=author,
            sections=sections,
            total_pages=len(sections),
            processed_pages=len(sections),
            char_count=sum(len(p.text) for s in sections for p in s.paragraphs),
        )

