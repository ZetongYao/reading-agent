from __future__ import annotations

import io
from pathlib import Path

import fitz
from docx import Document
from ebooklib import epub
from PIL import Image, ImageDraw

from backend.app.importers.docx import DOCXImporter
from backend.app.importers.epub import EPUBImporter
from backend.app.importers.pdf import PDFImporter
from backend.app.importers.plain_text import TextImporter


def progress(*_args):
    return None


class FakeOCR:
    def read_image(self, _image_bytes: bytes) -> str:
        return "A scanned page contains readable English text."


def test_text_pdf_extracts_every_page(tmp_path: Path):
    path = tmp_path / "text.pdf"
    document = fitz.open()
    for index in range(3):
        page = document.new_page()
        page.insert_text((72, 100), f"Page {index + 1} contains a complete English paragraph for testing.")
    document.set_metadata({"title": "Three Pages"})
    document.save(path)
    document.close()

    result = PDFImporter(ocr_service=FakeOCR()).import_book(path, progress)

    assert result.title == "Three Pages"
    assert result.total_pages == 3
    assert result.processed_pages == 3
    assert result.ocr_pages == 0
    assert len(result.sections) == 3
    assert "Page 3" in result.sections[2].paragraphs[0].text


def test_scanned_pdf_is_detected_and_ocr_continues(tmp_path: Path):
    image = Image.new("RGB", (900, 300), "white")
    ImageDraw.Draw(image).text((40, 100), "SCANNED ENGLISH PAGE", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    path = tmp_path / "scan.pdf"
    document = fitz.open()
    page = document.new_page(width=900, height=300)
    page.insert_image(page.rect, stream=buffer.getvalue())
    document.save(path)
    document.close()

    result = PDFImporter(ocr_service=FakeOCR()).import_book(path, progress)

    assert result.ocr_pages == 1
    assert result.failed_pages == []
    assert "scanned page" in result.sections[0].paragraphs[0].text


def test_epub_txt_and_docx_import(tmp_path: Path):
    text_path = tmp_path / "plain.txt"
    text_path.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")
    text_result = TextImporter().import_book(text_path, progress)
    assert [p.text for p in text_result.sections[0].paragraphs] == ["First paragraph.", "Second paragraph."]

    docx_path = tmp_path / "document.docx"
    document = Document()
    document.core_properties.title = "DOCX Book"
    document.add_heading("Opening", level=1)
    document.add_paragraph("The document has readable content.")
    document.save(docx_path)
    docx_result = DOCXImporter().import_book(docx_path, progress)
    assert docx_result.title == "DOCX Book"
    assert docx_result.sections[0].title == "Opening"

    epub_path = tmp_path / "book.epub"
    book = epub.EpubBook()
    book.set_identifier("test-book")
    book.set_title("EPUB Book")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter One", file_name="chapter.xhtml", lang="en")
    chapter.content = "<h1>Chapter One</h1><p>An EPUB sentence appears here.</p>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(str(epub_path), book)
    epub_result = EPUBImporter().import_book(epub_path, progress)
    assert epub_result.title == "EPUB Book"
    assert epub_result.sections[0].paragraphs[0].text == "An EPUB sentence appears here."

